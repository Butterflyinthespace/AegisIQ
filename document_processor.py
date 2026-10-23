"""AegisIQ: local evidence extraction for simulated documents."""
from __future__ import annotations
import csv
import io
import json
from pathlib import Path

MAX_BYTES = 10 * 1024 * 1024
SUPPORTED = {'.pdf', '.docx', '.xlsx', '.csv', '.json', '.txt'}
MAX_CHARS = 150_000


def extract_document(uploaded_file):
    """Return structured extraction result; never persist uploaded bytes."""
    name = Path(uploaded_file.name).name
    ext = Path(name).suffix.lower()
    data = uploaded_file.getvalue()
    if ext not in SUPPORTED:
        raise ValueError('Unsupported document format')
    if not data or len(data) > MAX_BYTES:
        raise ValueError('File must be non-empty and at most 10 MB')
    sections = []
    if ext == '.pdf':
        import fitz
        with fitz.open(stream=data, filetype='pdf') as doc:
            if doc.page_count > 150:
                raise ValueError('PDF exceeds 150 pages')
            for i, page in enumerate(doc):
                sections.append({'source': f'Page {i+1}', 'text': page.get_text('text')})
    elif ext == '.docx':
        from docx import Document
        document = Document(io.BytesIO(data))
        sections.append({'source': 'Paragraphs', 'text': '\n'.join(p.text for p in document.paragraphs if p.text.strip())})
        for i, table in enumerate(document.tables, 1):
            rows = [' | '.join(cell.text.replace('\n', ' ') for cell in row.cells) for row in table.rows]
            sections.append({'source': f'Table {i}', 'text': '\n'.join(rows)})
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in book.worksheets[:20]:
                lines = []
                for idx, row in enumerate(sheet.iter_rows(values_only=True)):
                    if idx >= 5000:
                        lines.append('[Additional rows omitted]')
                        break
                    cells = [str(v).replace('\n', ' ') if v is not None else '' for v in row[:50]]
                    if any(cells):
                        lines.append(' | '.join(cells))
                sections.append({'source': f'Sheet: {sheet.title}', 'text': '\n'.join(lines)})
        finally:
            book.close()
    elif ext == '.csv':
        content = data.decode('utf-8-sig')
        rows = csv.reader(io.StringIO(content))
        lines = []
        for idx, row in enumerate(rows):
            if idx >= 5000:
                lines.append('[Additional rows omitted]')
                break
            lines.append(' | '.join(row[:50]))
        sections.append({'source': 'CSV rows', 'text': '\n'.join(lines)})
    elif ext == '.json':
        obj = json.loads(data.decode('utf-8-sig'))
        sections.append({'source': 'JSON structure', 'text': json.dumps(obj, ensure_ascii=False, indent=2)})
    else:
        sections.append({'source': 'Text', 'text': data.decode('utf-8-sig')})
    total = sum(len(section['text']) for section in sections)
    remaining = MAX_CHARS
    for section in sections:
        original = section['text']
        section['text'] = original[:remaining]
        remaining -= len(section['text'])
    sections = [section for section in sections if section['text'].strip()]
    return {
        'filename': name,
        'format': ext[1:].upper(),
        'sections': sections,
        'characters': sum(len(s['text']) for s in sections),
        'truncated': total > MAX_CHARS,
        'empty': not sections,
    }
